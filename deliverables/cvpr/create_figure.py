from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
O=Path('/workspace/cvpr-figure');r=Presentation();r.slide_width=Pt(504);r.slide_height=Pt(248);s=r.slides.add_slide(r.slide_layouts[6])
ink='29384A';blue='DCEAF7';orange='F8DFC1';paper='F7F9FC';gray='E2E6EB';purple='D9D1EB';white='FFFFFF'
def box(x,y,w,h,fill=white,name='',dash=False,line=ink):
 a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h));a.name=name;a._element.spPr.append(OxmlElement('a:effectLst'));a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(fill);a.line.color.rgb=RGBColor.from_string(line);a.line.width=Pt(.55)
 if dash:a.line.dash_style=MSO_LINE_DASH_STYLE.DASH
 return a
def label(x,y,w,h,t,size=7.5,bold=False,align=PP_ALIGN.CENTER):
 a=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));a.name=t.replace('\n',' ');f=a.text_frame;f.clear();f.word_wrap=True;f.margin_left=f.margin_right=Pt(1);f.margin_top=f.margin_bottom=0;f.vertical_anchor=MSO_ANCHOR.MIDDLE
 for i,line in enumerate(t.split('\n')):
  p=f.paragraphs[0] if i==0 else f.add_paragraph();p.text=line;p.alignment=align;p.font.name='Liberation Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(ink);p.space_before=p.space_after=Pt(0)
 return a
def module(x,y,w,h,t,fill=blue,size=7.5):
 a=box(x,y,w,h,fill,t.replace('\n',' '));label(x+1,y+2,w-2,h-4,t,size);return a
def seg(x,y,X,Y,head=True,both=False):
 a=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x),Pt(y),Pt(X),Pt(Y));a._element.spPr.append(OxmlElement('a:effectLst'));a.line.width=Pt(.65);a.line.color.rgb=RGBColor.from_string(ink)
 ln=a._element.spPr.get_or_add_ln()
 if head:
  e=OxmlElement('a:tailEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');ln.append(e)
 if both:
  e=OxmlElement('a:headEnd');e.set('type','triangle');e.set('w','sm');e.set('len','sm');ln.append(e)
 return a
def route(points):
 for i,(a,b) in enumerate(zip(points,points[1:])):seg(*a,*b,head=i==len(points)-2)
def grid(x,y,cols,rows,sz=5,gap=2,colors=None):
 colors=colors or [blue]
 for j in range(rows):
  for i in range(cols):box(x+i*(sz+gap),y+j*(sz+gap),sz,sz,colors[(i+j*cols)%len(colors)],'Editable token',line='75879C')
def group(x,y,w,h,t):
 box(x,y,w,h,paper,t,True,'8697AA');label(x+2,y+3,w-4,14,t,8,True)
# Stage groups at final print size; no in-figure paper title.
group(2,15,100,202,'Image encoding')
group(108,15,65,202,'Dual branches')
group(179,15,104,202,'Text-guided screening')
group(289,15,66,202,'Interaction')
group(361,15,72,202,'Fusion & refinement')
# Image input and window tiles.
s.shapes.add_picture('/workspace/reconstruction/input_illustration.png',Pt(9),Pt(48),width=Pt(38),height=Pt(42));label(7,91,42,17,'Input image',7)
for dx,dy in [(0,0),(3,-3),(6,-6)]:
 box(20+dx,125+dy,13,17,white,'Sliding-window tile');box(22+dx,127+dy,9,6,blue,'Tile feature')
label(6,149,46,32,'Sliding Window\nPartition',7)
module(57,99,39,60,'Shared\nCLIP\nLocal\nEncoder',blue,7)
seg(28,110,28,119);seg(39,132,57,132)
# Parallel branches.
box(113,48,55,62,white,'Local branch');label(114,51,53,24,'Local Branch\nExtraction',7.5);grid(121,82,5,2,5,2,[blue,blue,orange,gray])
box(113,142,55,62,white,'Stitch global branch');label(114,145,53,24,'Stitch Global\nStitch Branch',7.5);grid(121,177,5,2,5,2,[blue,gray,blue,blue])
route([(96,103),(105,103),(105,79),(113,79)])
route([(96,141),(105,141),(105,173),(113,173)])
# Three independent incoming modalities into proxy anchors.
module(188,41,86,31,'Offline CBP\nText Prototype',white)
module(188,94,86,35,'Text-Guided Proxy\nAnchor Generation',orange)
seg(231,72,231,94)
route([(168,79),(177,79),(177,105),(188,105)])
route([(168,173),(182,173),(182,118),(188,118)])
module(188,153,86,29,'Lightweight\nToken Selection',blue)
seg(231,129,231,153)
grid(205,192,8,2,5,2,[blue,gray,orange,gray,purple,gray,blue,gray]);label(189,209,84,7,'Selected tokens',6)
# Two interacting attention modules; no spurious feedback from fusion.
module(295,80,54,32,'Local\nAttention',blue)
module(295,153,54,37,'Global\nHierarchical\nAttention',blue,7)
route([(274,161),(286,161),(286,96),(295,96)])
seg(274,172,295,172)
seg(322,112,322,153,both=True)
# Attention outputs feed a single fusion module.
module(367,96,60,39,'Scale-Gated\nFeature Fusion',orange)
route([(349,96),(358,96),(358,106),(367,106)])
route([(349,172),(358,172),(358,125),(367,125)])
module(367,164,60,39,'Cross-Modal\nConsistency\nRefinement',blue,7.2)
seg(397,135,397,164)
# Output image is intentionally a raster illustration, not a claimed result.
s.shapes.add_picture('/workspace/reconstruction/output_illustration.png',Pt(443),Pt(161),width=Pt(55),height=Pt(43));box(443,161,55,43,white,'Output image border').fill.background()
seg(427,183,443,183);label(438,207,65,26,'Output\nSegmentation Map',7.5)
# Short native legend: colors support rather than replace text labels.
for x,c,t in [(75,blue,'Visual features'),(195,orange,'Proxy / fusion'),(310,purple,'Text prototypes')]:
 box(x,232,6,6,c,'Legend swatch');label(x+9,229,100,12,t,6.5,align=PP_ALIGN.LEFT)
for shape in s.shapes:
 for effect in shape._element.xpath('.//a:effectRef'):
  effect.set('idx','0')
r.save(O/'TS_CLIP_CVPR_framework.pptx')
(O/'caption_and_usage.txt').write_text('Figure. Overview of TS-CLIP. A shared CLIP local encoder produces local and stitched-global features. Together with offline CBP text prototypes, these features guide proxy-anchor generation and lightweight token selection. Local and global hierarchical attention exchange information before scale-gated fusion and cross-modal consistency refinement produce the segmentation map.\n\nUse the PDF as a two-column-wide figure at width=\\textwidth in the CVPR template (figure* environment). The canvas is 7.0 x 3.444 inches. Labels are 7–8 pt, with a 6 pt token annotation and 6.5 pt legend. The artwork is a layout draft, not a certification of a particular year\'s submission requirements.\n\nInput/output thumbnails are illustrative bitmap crops from an earlier generated schematic, not measured experimental outputs. Other elements are native editable PowerPoint objects.\n\nAll original pipeline stages are preserved; similarity matching and top-k internals shown in the reference are not added as separately claimed stages.\n')
print('Created native PPTX:',len(s.shapes),'objects')
